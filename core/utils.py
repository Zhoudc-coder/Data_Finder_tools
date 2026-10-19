# -*- coding: utf-8 -*-
"""
通用工具函数：
- 跨平台路径处理
- 索引库/日志目录回退
- CSV 编码兼容读取
- 输入解析、文件名过滤、时间格式化
- 从 Excel/CSV 导入第一列（SN 或过程码）
- 网络共享盘（UNC/SMB）友好的目录扫描与路径规范化

【重要安全约定】
本模块内所有读取原始数据文件（“待出库”文件夹内的 .csv）的函数，
均严格使用 "r" 只读模式打开，或 openpyxl 的 read_only=True 模式读取，
不对原始数据文件做任何 "w"/"a"/"x"/"r+" 写入或修改，也不会保存回原文件。
"""
from __future__ import annotations

import csv
import logging
import os
import platform
import re
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Callable, Optional

APP_NAME = "LogisticsDataFinder"

# CSV 编码尝试顺序
CSV_ENCODINGS = ("utf-8-sig", "gbk", "gb18030")

# 导入 SN 时视为标题行的常见值（不区分大小写）
SN_HEADER_VALUES = {"sn", "s/n", "序列号", "sn号", "sn码", "serial", "产品序列号"}

# 导入过程码时视为标题行的常见值（不区分大小写）
PROCESS_HEADER_VALUES = {
    "过程码", "流程码", "过程编号", "过程号", "文件名", "文件名称",
    "process", "process_code", "processcode", "code",
}


class CancelledError(Exception):
    """用户取消操作。"""
    pass


# ----------------------------------------------------------------------
# 跨平台路径与网络路径判断（均不触发网络 IO）
# ----------------------------------------------------------------------
def is_network_path(path) -> bool:
    """
    判断是否为网络路径：
    - UNC 路径：\\\\server\\share\\... 或 //server/share/...
    - Windows 映射网络驱动器：Z:\\...
    仅做字符串与本地 API 判断，不触发网络 IO，可在 UI 线程安全调用。
    """
    try:
        s = str(path).strip()
    except Exception:
        return False
    if not s:
        return False
    if s.startswith("\\\\") or s.startswith("//"):
        return True
    if os.name == "nt":
        try:
            import ctypes

            drive = os.path.splitdrive(s)[0]
            if drive and len(drive) == 2 and drive[1] == ":":
                DRIVE_REMOTE = 4
                if ctypes.windll.kernel32.GetDriveTypeW(drive + "\\") == DRIVE_REMOTE:
                    return True
        except Exception:
            pass
    return False


def normalize_path_str(path) -> str:
    """
    规范化路径字符串，用作索引库主键。
    仅做字符串级规范化（abspath + normpath），不解析符号链接，
    避免在网络共享盘（UNC/SMB）上触发昂贵的网络查询。
    """
    try:
        return os.path.normpath(os.path.abspath(str(path)))
    except Exception:
        return str(path)


def is_path_inside(child, parent) -> bool:
    """
    判断 child 是否位于 parent 目录内（含子目录）。
    使用字符串级规范化，不解析符号链接，不触发网络 IO。
    Windows 下大小写不敏感。
    """
    try:
        child_norm = normalize_path_str(child)
        parent_norm = normalize_path_str(parent)
        if os.name == "nt":
            child_norm = child_norm.casefold()
            parent_norm = parent_norm.casefold()
        if child_norm == parent_norm:
            return True
        sep = "\\" if os.name == "nt" else "/"
        if not parent_norm.endswith(sep):
            parent_norm += sep
        return child_norm.startswith(parent_norm)
    except Exception:
        return False


def get_app_dir() -> Path:
    """
    获取程序目录。
    - 源码运行：项目根目录
    - PyInstaller 打包后：exe 所在目录
    使用 abspath 而非 resolve，避免网络盘上的额外查询。
    """
    if getattr(sys, "frozen", False):
        return Path(os.path.dirname(os.path.abspath(sys.executable)))
    return Path(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


def get_user_data_dir() -> Path:
    """
    获取跨平台用户可写目录。
    Windows: %LOCALAPPDATA%/LogisticsDataFinder
    macOS: ~/Library/Application Support/LogisticsDataFinder
    Linux: $XDG_DATA_HOME/LogisticsDataFinder 或 ~/.local/share/LogisticsDataFinder
    """
    system = platform.system()
    if system == "Windows":
        base = os.environ.get("LOCALAPPDATA") or os.environ.get("APPDATA")
        if base:
            path = Path(base) / APP_NAME
        else:
            path = Path.home() / "AppData" / "Local" / APP_NAME
    elif system == "Darwin":
        path = Path.home() / "Library" / "Application Support" / APP_NAME
    else:
        base = os.environ.get("XDG_DATA_HOME")
        path = Path(base) / APP_NAME if base else Path.home() / ".local" / "share" / APP_NAME

    path.mkdir(parents=True, exist_ok=True)
    return path


def _is_writable_dir(path: Path) -> bool:
    """测试目录是否可写。"""
    try:
        path.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=path, prefix=".write_test_", delete=True):
            pass
        return True
    except Exception:
        return False


def get_index_db_path() -> tuple[Path, str]:
    """
    获取 SQLite 索引库路径。
    优先程序目录；程序目录不可写时自动回退到用户可写目录。
    若程序目录本身位于网络路径，直接使用本地用户目录，避免 SQLite 在 SMB 上的兼容性问题。
    """
    app_dir = get_app_dir()
    if is_network_path(app_dir):
        fallback = get_user_data_dir() / "index_cache.db"
        return fallback, f"程序目录位于网络路径，索引库已放到本地用户目录：{fallback}"

    candidate = app_dir / "index_cache.db"
    if _is_writable_dir(app_dir):
        return candidate, f"索引库位置：{candidate}（程序目录可写）"

    fallback = get_user_data_dir() / "index_cache.db"
    return fallback, f"程序目录不可写，索引库已自动回退到：{fallback}"


def get_log_dir() -> Path:
    """日志目录，始终使用用户可写目录。"""
    path = get_user_data_dir() / "logs"
    path.mkdir(parents=True, exist_ok=True)
    return path


def setup_logging() -> Path:
    """初始化日志，写入用户目录，避免打包后程序目录不可写。"""
    log_file = get_log_dir() / "app.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        handlers=[
            logging.FileHandler(log_file, encoding="utf-8"),
            logging.StreamHandler(),
        ],
        force=True,
    )
    logging.info("程序启动，日志文件：%s", log_file)
    return log_file


def sanitize_filename(name: str, max_len: int = 120) -> str:
    """过滤 Windows 非法文件名字符。"""
    name = re.sub(r'[\\/:*?"<>|]+', "_", str(name))
    name = re.sub(r"\s+", "_", name.strip())
    if not name:
        name = "export"
    return name[:max_len]


def format_mtime(mtime: float) -> str:
    """文件修改时间格式化。"""
    try:
        return datetime.fromtimestamp(mtime).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return ""


def parse_input_tokens(text: str, unique: bool = True) -> list[str]:
    """
    解析用户输入。
    支持多行、英文逗号、中文逗号、空格、分号分隔。
    unique=True 时保留首次出现顺序去重。
    """
    tokens = [t for t in re.split(r"[\s,，;；]+", text or "") if t]
    if not unique:
        return tokens

    seen = set()
    result = []
    for token in tokens:
        if token not in seen:
            seen.add(token)
            result.append(token)
    return result


def _set_csv_field_limit() -> None:
    """允许较大单元格，避免 csv.Error: field larger than field limit。"""
    limit = 1024 * 1024 * 1024
    while True:
        try:
            csv.field_size_limit(limit)
            break
        except OverflowError:
            limit //= 10


_set_csv_field_limit()


# ----------------------------------------------------------------------
# 网络共享盘友好的目录扫描
# ----------------------------------------------------------------------
def scan_csv_files(
    folder,
    progress_cb: Optional[Callable[[int, int, str], None]] = None,
    log_cb: Optional[Callable[[str], None]] = None,
    cancel_event=None,
) -> list[Path]:
    """
    递归扫描文件夹下所有 .csv 文件。
    使用 os.scandir，利用 DirEntry 缓存的文件类型属性，
    避免对每个条目额外 stat，适合网络共享盘（UNC/SMB）。

    progress_cb(dirs_scanned, found_count, message)：
      - dirs_scanned：已遍历目录数（非总目录数，网络盘无法预知总数）
      - found_count：已发现 CSV 文件数
    """
    folder = Path(folder)
    result: list[Path] = []
    stack: list[Path] = [folder]
    dirs_scanned = 0
    last_report = 0

    while stack:
        if cancel_event is not None and cancel_event.is_set():
            break
        current = stack.pop()
        dirs_scanned += 1
        try:
            with os.scandir(current) as it:
                for idx, entry in enumerate(it):
                    if cancel_event is not None and cancel_event.is_set() and idx % 100 == 0:
                        return result
                    try:
                        if entry.is_dir(follow_symlinks=False):
                            stack.append(Path(entry.path))
                        elif entry.is_file():
                            if entry.name.lower().endswith(".csv"):
                                result.append(Path(entry.path))
                    except OSError as e:
                        if log_cb:
                            log_cb(f"跳过无法访问的条目：{entry.path}，原因：{e}")
        except OSError as e:
            if log_cb:
                log_cb(f"无法列出目录：{current}，原因：{e}")
            continue

        if progress_cb and (dirs_scanned - last_report >= 10):
            last_report = dirs_scanned
            progress_cb(
                dirs_scanned,
                len(result),
                f"正在遍历目录：已扫描 {dirs_scanned} 个目录，发现 {len(result)} 个 CSV 文件",
            )

    if progress_cb:
        progress_cb(
            dirs_scanned,
            len(result),
            f"目录遍历完成：共扫描 {dirs_scanned} 个目录，发现 {len(result)} 个 CSV 文件",
        )
    return result


# ======================================================================
# 以下所有读取原始数据文件的函数均为只读访问，不会对源文件产生任何写入。
# ======================================================================
def read_sn_column(file_path: Path, log_cb: Optional[Callable[[str], None]] = None) -> list[tuple[int, str]]:
    """
    【只读】读取 CSV 第一列 SN。
    返回 [(数据行号0-based, SN), ...]。
    第一行视为标题行；数据行号从 0 开始，不含标题行。
    严格使用 "r" 模式打开，不写入、不修改原文件。
    遇到 OSError（如网络断连）直接放弃，不再换编码重试。
    """
    last_err = None

    for encoding in CSV_ENCODINGS:
        try:
            with open(file_path, "r", encoding=encoding, newline="") as f:
                reader = csv.reader(f)
                try:
                    next(reader)  # 跳过标题行
                except StopIteration:
                    if log_cb:
                        log_cb(f"空文件：{file_path}")
                    return []

                result: list[tuple[int, str]] = []
                for data_idx, row in enumerate(reader):
                    if not row:
                        continue
                    sn = row[0].strip()
                    if sn:
                        result.append((data_idx, sn))
                return result

        except UnicodeDecodeError as e:
            last_err = e
            continue
        except csv.Error as e:
            last_err = e
            continue
        except OSError as e:
            # 网络/磁盘 IO 错误，换编码也无用，直接放弃
            last_err = e
            break
        except Exception as e:
            last_err = e
            break

    if log_cb:
        log_cb(f"读取 SN 列失败：{file_path}，原因：{last_err}")
    return []


def read_csv_rows_by_indices(
    file_path: Path,
    indices: set[int],
    log_cb: Optional[Callable[[str], None]] = None,
) -> tuple[list[str], dict[int, list[str]]]:
    """
    【只读】读取指定数据行号的完整行。
    返回 (标题行, {数据行号0-based: 行数据})。
    严格使用 "r" 模式打开，不写入、不修改原文件。
    """
    wanted = set(indices)
    last_err = None

    for encoding in CSV_ENCODINGS:
        try:
            with open(file_path, "r", encoding=encoding, newline="") as f:
                reader = csv.reader(f)
                try:
                    header = next(reader)
                except StopIteration:
                    return [], {}

                if not wanted:
                    return header, {}

                result: dict[int, list[str]] = {}
                for data_idx, row in enumerate(reader):
                    if data_idx in wanted:
                        result[data_idx] = row
                        if len(result) == len(wanted):
                            break
                return header, result

        except UnicodeDecodeError as e:
            last_err = e
            continue
        except csv.Error as e:
            last_err = e
            continue
        except OSError as e:
            last_err = e
            break
        except Exception as e:
            last_err = e
            break

    if log_cb:
        log_cb(f"读取完整行失败：{file_path}，原因：{last_err}")
    return [], {}


def read_all_csv_rows(
    file_path: Path,
    log_cb: Optional[Callable[[str], None]] = None,
) -> tuple[list[str], list[list[str]]]:
    """
    【只读】读取 CSV 全部数据行（不含标题行）。
    返回 (标题行, 数据行列表)。
    严格使用 "r" 模式打开，不写入、不修改原文件。
    """
    last_err = None

    for encoding in CSV_ENCODINGS:
        try:
            with open(file_path, "r", encoding=encoding, newline="") as f:
                reader = csv.reader(f)
                try:
                    header = next(reader)
                except StopIteration:
                    return [], []

                rows = [row for row in reader if row]
                return header, rows

        except UnicodeDecodeError as e:
            last_err = e
            continue
        except csv.Error as e:
            last_err = e
            continue
        except OSError as e:
            last_err = e
            break
        except Exception as e:
            last_err = e
            break

    if log_cb:
        log_cb(f"读取文件全部数据失败：{file_path}，原因：{last_err}")
    return [], []


# ----------------------------------------------------------------------
# 从文件第一列导入 SN / 过程码（只读）
# ----------------------------------------------------------------------
def load_first_column_from_file(
    path: Path,
    header_values: Optional[set[str]] = None,
    log_cb: Optional[Callable[[str], None]] = None,
) -> list[str]:
    """
    【只读】从 .xlsx 或 .csv 文件第一列读取待查找值（SN 或过程码）。
    兼容有/无标题行；自动去重保序。
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"文件不存在：{path}")
    if not path.is_file():
        raise ValueError(f"不是文件：{path}")

    if header_values is None:
        header_values = SN_HEADER_VALUES
    normalized_headers = {str(h).strip().lower() for h in header_values}

    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        values = _load_first_column_from_xlsx(path, log_cb)
    elif suffix == ".csv":
        values = _load_first_column_from_csv(path, log_cb)
    elif suffix in (".xls", ".xlsm"):
        raise ValueError(f"暂不支持 {suffix}，请另存为 .xlsx 后再导入")
    else:
        raise ValueError(f"不支持的文件格式：{suffix}，请选择 .xlsx 或 .csv 文件")

    cleaned: list[str] = []
    for v in values:
        s = str(v).strip()
        if s:
            cleaned.append(s)

    if cleaned and cleaned[0].lower() in normalized_headers:
        cleaned = cleaned[1:]

    seen = set()
    result: list[str] = []
    for v in cleaned:
        if v not in seen:
            seen.add(v)
            result.append(v)
    return result


def _load_first_column_from_xlsx(path: Path, log_cb=None) -> list[str]:
    """【只读】使用 openpyxl 流式读取 xlsx 第一列。"""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ValueError("缺少 openpyxl 依赖，请先安装：pip install openpyxl") from e

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        values: list[str] = []
        for row in ws.iter_rows(values_only=True, max_col=1):
            if not row:
                continue
            v = row[0]
            if v is None:
                continue
            if isinstance(v, float) and v.is_integer():
                values.append(str(int(v)))
            else:
                s = str(v).strip()
                if s:
                    values.append(s)
        return values
    except Exception as e:
        if log_cb:
            log_cb(f"读取 xlsx 失败：{path}，原因：{e}")
        raise
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _load_first_column_from_csv(path: Path, log_cb=None) -> list[str]:
    """【只读】按 CSV_ENCODINGS 顺序读取 csv 第一列。"""
    last_err = None
    for encoding in CSV_ENCODINGS:
        try:
            with open(path, "r", encoding=encoding, newline="") as f:
                reader = csv.reader(f)
                values: list[str] = []
                for row in reader:
                    if row and row[0]:
                        values.append(row[0])
                return values
        except UnicodeDecodeError as e:
            last_err = e
            continue
        except csv.Error as e:
            last_err = e
            continue
        except OSError as e:
            last_err = e
            break
        except Exception as e:
            last_err = e
            break

    if log_cb:
        log_cb(f"读取 CSV 失败：{path}，原因：{last_err}")
    raise ValueError(f"无法读取 CSV 文件：{path}，原因：{last_err}")


# 兼容旧名字
load_sn_from_file = load_first_column_from_file